# admin.py — Admin Dashboard
# Features: manage quizzes, questions, view attempts, student progress,
#            import from Excel, export to Excel.

import tkinter as tk
from tkinter import messagebox, filedialog
import db

# Try to import openpyxl for Excel import/export
try:
    import openpyxl
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

import main

# ── Colors & Fonts ────────────────────────────────────────
BG      = "#1e1e2e"
PANEL   = "#2a2a3e"
ACCENT  = "#7c3aed"
ENTRY   = "#12121c"
FG      = "#ffffff"
DIM     = "#c0c0d0"
GREEN   = "#22c55e"
RED     = "#ef4444"

FONT_TITLE = ("Segoe UI", 18, "bold")
FONT_BODY  = ("Segoe UI", 11)
FONT_SMALL = ("Segoe UI", 10)
FONT_TINY  = ("Segoe UI", 9)


# ── Helper Functions ───────────────────────────────────────

def make_btn(parent, text, command, color=ACCENT):
    """
    Creates a colored button (Frame + Label trick for macOS compatibility).
    Clicking either the frame or the label triggers command().
    """
    frame = tk.Frame(parent, bg=color, cursor="hand2")
    label = tk.Label(frame, text=text, bg=color, fg=FG,
                     font=FONT_BODY, padx=10, pady=7)
    label.pack()
    frame.bind("<Button-1>", lambda e: command())
    label.bind("<Button-1>", lambda e: command())
    return frame


def make_field(parent, label_text, value=""):
    """Creates a label + entry field and returns the StringVar."""
    tk.Label(parent, text=label_text, bg=PANEL, fg=DIM,
             font=FONT_TINY).pack(anchor="w", pady=(8, 2))
    var = tk.StringVar(value=value)
    tk.Entry(parent, textvariable=var, bg=ENTRY, fg=FG,
             relief="flat", font=FONT_BODY, insertbackground=FG).pack(
             fill="x", ipady=6)
    return var


def center_window(win, width, height):
    """Centers a window on the screen."""
    win.update_idletasks()
    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()
    x  = (sw - width)  // 2
    y  = (sh - height) // 2
    win.geometry(f"{width}x{height}+{x}+{y}")


# ── Admin Dashboard ────────────────────────────────────────

def show_admin(user_id, username):
    """Main admin window: quiz list + all action buttons."""

    root = tk.Tk()
    root.title("Admin Dashboard")
    root.configure(bg=BG)
    center_window(root, 900, 620)

    # ── Top bar ───────────────────────────────────────────
    top_bar = tk.Frame(root, bg=ACCENT, pady=10)
    top_bar.pack(fill="x")
    tk.Label(top_bar, text=f"🛡  Admin Panel  —  {username}",
             font=FONT_TITLE, bg=ACCENT, fg=FG).pack(side="left", padx=20)

    def do_logout():
        if messagebox.askyesno("Logout", "Are you sure you want to logout?"):
            root.destroy()
            main.show_login()

    make_btn(top_bar, "⇠ Logout", do_logout, color="#5b21b6").pack(
        side="right", padx=16, pady=2)

    # ── Body ──────────────────────────────────────────────
    body = tk.Frame(root, bg=BG)
    body.pack(fill="both", expand=True, padx=20, pady=12)

    tk.Label(body, text="Quiz List",
             font=("Segoe UI", 13, "bold"), bg=BG, fg=FG).pack(anchor="w")

    # Quiz listbox with scrollbar
    list_frame = tk.Frame(body, bg=ENTRY)
    list_frame.pack(fill="both", expand=True, pady=(4, 6))

    list_scroll = tk.Scrollbar(list_frame)
    list_scroll.pack(side="right", fill="y")

    quiz_listbox = tk.Listbox(list_frame, bg=ENTRY, fg=FG, font=FONT_BODY,
                              relief="flat", selectbackground=ACCENT,
                              activestyle="none", yscrollcommand=list_scroll.set)
    quiz_listbox.pack(fill="both", expand=True)
    list_scroll.config(command=quiz_listbox.yview)

    all_quizzes = []   # quizzes fetched from the database

    def load_quizzes():
        """Fetch all quizzes from DB and display in the listbox."""
        all_quizzes.clear()
        all_quizzes.extend(db.get_all_quizzes())
        quiz_listbox.delete(0, tk.END)
        for quiz in all_quizzes:
            quiz_listbox.insert(tk.END, f"  {quiz[1]}   ⏱ {quiz[2]} min")

    load_quizzes()

    def get_selected_quiz():
        """Return the quiz tuple for the currently selected row, or None."""
        sel = quiz_listbox.curselection()
        if not sel:
            messagebox.showwarning("No selection", "Please click a quiz first.")
            return None
        return all_quizzes[sel[0]]

    # ── Button rows ───────────────────────────────────────
    # Row 1: quiz management
    row1 = tk.Frame(body, bg=BG)
    row1.pack(fill="x", pady=(0, 4))

    # Row 2: reporting + Excel
    row2 = tk.Frame(body, bg=BG)
    row2.pack(fill="x")

    # ── Callbacks ─────────────────────────────────────────

    def on_add_quiz():
        open_add_quiz_window(root, on_done=load_quizzes)

    def on_edit_quiz():
        quiz = get_selected_quiz()
        if quiz:
            open_edit_quiz_window(root, quiz, on_done=load_quizzes)

    def on_add_question():
        quiz = get_selected_quiz()
        if quiz:
            open_add_question_window(root, quiz[0], quiz[1])

    def on_edit_questions():
        quiz = get_selected_quiz()
        if quiz:
            open_edit_questions_window(root, quiz[0], quiz[1])

    def on_delete_quiz():
        quiz = get_selected_quiz()
        if not quiz:
            return
        confirmed = messagebox.askyesno(
            "Delete Quiz",
            f'Delete "{quiz[1]}" and ALL its questions and attempts?\n\nThis cannot be undone!'
        )
        if confirmed:
            db.delete_quiz(quiz[0])
            load_quizzes()

    def on_all_attempts():
        open_all_attempts_window(root)

    def on_student_progress():
        open_student_progress_window(root)

    def on_import_excel():
        open_excel_import_window(root, on_done=load_quizzes)

    def on_export_excel():
        export_gradebook_to_excel(root)

    # Add buttons to row 1
    for text, cmd, color in [
        ("➕ Add Quiz",       on_add_quiz,       ACCENT),
        ("📝 Edit Quiz",      on_edit_quiz,       "#0ea5e9"),
        ("❓ Add Question",   on_add_question,    ACCENT),
        ("✏️ Edit Questions", on_edit_questions,  "#0ea5e9"),
        ("🗑 Delete Quiz",    on_delete_quiz,     RED),
    ]:
        make_btn(row1, text, cmd, color=color).pack(side="left", padx=(0, 5))

    # Add buttons to row 2
    for text, cmd, color in [
        ("📊 All Attempts",     on_all_attempts,     ACCENT),
        ("👤 Student Progress", on_student_progress, "#7e22ce"),
        ("📥 Import Excel",     on_import_excel,     "#16a34a"),
        ("📤 Export Excel",     on_export_excel,     "#0891b2"),
    ]:
        make_btn(row2, text, cmd, color=color).pack(side="left", padx=(0, 5))

    root.mainloop()


# ── Add Quiz Window ────────────────────────────────────────

def open_add_quiz_window(parent, on_done):
    """Small form to create a new quiz."""
    win = tk.Toplevel(parent)
    win.title("Add New Quiz")
    win.configure(bg=BG)
    center_window(win, 380, 240)
    win.grab_set()

    tk.Label(win, text="Add New Quiz", font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(20, 8))

    panel = tk.Frame(win, bg=PANEL, padx=20, pady=16)
    panel.pack(padx=20, fill="both", expand=True)

    title_var = make_field(panel, "Quiz Title")
    time_var  = make_field(panel, "Time Limit (minutes)", value="30")

    error_var = tk.StringVar()
    tk.Label(panel, textvariable=error_var, bg=PANEL, fg="#f87171",
             font=FONT_TINY).pack(pady=(4, 0))

    def save():
        title = title_var.get().strip()
        time  = time_var.get().strip()
        if not title:
            error_var.set("Title is required."); return
        if not time.isdigit() or int(time) < 1:
            error_var.set("Enter a valid number of minutes."); return
        db.add_quiz(title, int(time))
        on_done()
        win.destroy()

    make_btn(panel, "Create Quiz", save, color=GREEN).pack(pady=(8, 0), fill="x")


# ── Edit Quiz Window ───────────────────────────────────────

def open_edit_quiz_window(parent, quiz, on_done):
    """Pre-filled form to rename a quiz or change its time limit."""
    quiz_id, old_title, old_time = quiz

    win = tk.Toplevel(parent)
    win.title("Edit Quiz")
    win.configure(bg=BG)
    center_window(win, 380, 240)
    win.grab_set()

    tk.Label(win, text="Edit Quiz", font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(20, 8))

    panel = tk.Frame(win, bg=PANEL, padx=20, pady=16)
    panel.pack(padx=20, fill="both", expand=True)

    title_var = make_field(panel, "Quiz Title",           value=old_title)
    time_var  = make_field(panel, "Time Limit (minutes)", value=str(old_time))

    error_var = tk.StringVar()
    tk.Label(panel, textvariable=error_var, bg=PANEL, fg="#f87171",
             font=FONT_TINY).pack(pady=(4, 0))

    def save():
        title = title_var.get().strip()
        time  = time_var.get().strip()
        if not title:
            error_var.set("Title is required."); return
        if not time.isdigit() or int(time) < 1:
            error_var.set("Enter a valid number of minutes."); return
        db.update_quiz(quiz_id, title, int(time))
        on_done()
        win.destroy()
        messagebox.showinfo("Saved", f'Quiz updated to:\n"{title}"  —  {time} min')

    make_btn(panel, "Save Changes", save, color=GREEN).pack(pady=(8, 0), fill="x")


# ── Add Question Window ────────────────────────────────────

def open_add_question_window(parent, quiz_id, quiz_title):
    """Form to add a new question with 4 options to a quiz."""
    win = tk.Toplevel(parent)
    win.title(f"Add Question — {quiz_title}")
    win.configure(bg=BG)
    center_window(win, 480, 520)
    win.grab_set()

    tk.Label(win, text="Add Question", font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(16, 2))
    tk.Label(win, text=quiz_title, font=FONT_BODY, bg=BG, fg=DIM).pack()

    # Scrollable form
    canvas = tk.Canvas(win, bg=BG, highlightthickness=0)
    scrollbar = tk.Scrollbar(win, command=canvas.yview)
    canvas.configure(yscrollcommand=scrollbar.set)
    scrollbar.pack(side="right", fill="y")
    canvas.pack(fill="both", expand=True)

    form = tk.Frame(canvas, bg=PANEL, padx=20, pady=16)
    canvas.create_window((0, 0), window=form, anchor="nw")

    def update_scroll(e):
        canvas.configure(scrollregion=canvas.bbox("all"))
    form.bind("<Configure>", update_scroll)

    q_text_var = make_field(form, "Question Text")
    marks_var  = make_field(form, "Marks", value="1")

    tk.Label(form, text="Options (select the correct one):",
             bg=PANEL, fg=DIM, font=FONT_TINY).pack(anchor="w", pady=(10, 2))

    option_vars  = []   # StringVar for each option text
    correct_var  = tk.IntVar(value=0)   # index of the correct option (0-3)

    for i, letter in enumerate(["A", "B", "C", "D"]):
        row = tk.Frame(form, bg=PANEL)
        row.pack(fill="x", pady=2)
        tk.Radiobutton(row, text=letter, variable=correct_var, value=i,
                       bg=PANEL, fg=FG, selectcolor=ACCENT,
                       activebackground=PANEL).pack(side="left")
        var = tk.StringVar()
        tk.Entry(row, textvariable=var, bg=ENTRY, fg=FG, relief="flat",
                 font=FONT_BODY, insertbackground=FG).pack(
                 side="left", fill="x", expand=True, ipady=5, padx=4)
        option_vars.append(var)

    error_var = tk.StringVar()
    tk.Label(form, textvariable=error_var, bg=PANEL, fg="#f87171",
             font=FONT_TINY).pack(pady=(4, 0))

    def save():
        q_text = q_text_var.get().strip()
        marks  = marks_var.get().strip()
        opts   = [v.get().strip() for v in option_vars]

        if not q_text:
            error_var.set("Question text is required."); return
        if not marks.isdigit() or int(marks) < 1:
            error_var.set("Enter valid marks (number ≥ 1)."); return
        if any(o == "" for o in opts):
            error_var.set("All 4 options must be filled."); return

        q_id = db.add_question(quiz_id, q_text, int(marks))
        correct_index = correct_var.get()
        for i, opt_text in enumerate(opts):
            is_correct = 1 if i == correct_index else 0
            db.add_option(q_id, opt_text, is_correct)

        messagebox.showinfo("Added", "Question added successfully!")
        win.destroy()

    make_btn(form, "Save Question", save, color=GREEN).pack(pady=(12, 0), fill="x")


# ── Edit Questions Window ──────────────────────────────────

def open_edit_questions_window(parent, quiz_id, quiz_title):
    """Lists all questions for a quiz; click one to edit it."""
    win = tk.Toplevel(parent)
    win.title(f"Edit Questions — {quiz_title}")
    win.configure(bg=BG)
    center_window(win, 560, 480)

    tk.Label(win, text=f"Questions — {quiz_title}",
             font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(16, 4))

    q_list_frame  = tk.Frame(win, bg=ENTRY)
    q_list_frame.pack(fill="both", expand=True, padx=20, pady=6)

    q_scroll = tk.Scrollbar(q_list_frame)
    q_scroll.pack(side="right", fill="y")

    q_listbox = tk.Listbox(q_list_frame, bg=ENTRY, fg=FG, font=FONT_BODY,
                           relief="flat", selectbackground=ACCENT,
                           activestyle="none", yscrollcommand=q_scroll.set)
    q_listbox.pack(fill="both", expand=True)
    q_scroll.config(command=q_listbox.yview)

    questions = []

    def load():
        questions.clear()
        questions.extend(db.get_questions(quiz_id))
        q_listbox.delete(0, tk.END)
        for i, q in enumerate(questions):
            q_listbox.insert(tk.END, f"  Q{i+1}: {q[2][:60]}...")

    load()

    def on_edit():
        sel = q_listbox.curselection()
        if not sel:
            messagebox.showwarning("Select", "Click a question first.")
            return
        open_edit_single_question(win, questions[sel[0]], on_done=load)

    make_btn(win, "✏️ Edit Selected Question", on_edit, color="#0ea5e9").pack(pady=8)


# ── Edit Single Question ───────────────────────────────────

def open_edit_single_question(parent, question, on_done):
    """Form to edit one question's text, marks, options, and correct answer."""
    q_id, _, q_text, marks = question

    win = tk.Toplevel(parent)
    win.title("Edit Question")
    win.configure(bg=BG)
    center_window(win, 480, 520)
    win.grab_set()

    canvas = tk.Canvas(win, bg=BG, highlightthickness=0)
    scrollbar = tk.Scrollbar(win, command=canvas.yview)
    canvas.configure(yscrollcommand=scrollbar.set)
    scrollbar.pack(side="right", fill="y")
    canvas.pack(fill="both", expand=True)

    form = tk.Frame(canvas, bg=PANEL, padx=20, pady=16)
    canvas.create_window((0, 0), window=form, anchor="nw")
    form.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))

    q_text_var = make_field(form, "Question Text", value=q_text)
    marks_var  = make_field(form, "Marks",         value=str(marks))

    tk.Label(form, text="Options (select the correct one):",
             bg=PANEL, fg=DIM, font=FONT_TINY).pack(anchor="w", pady=(10, 2))

    options      = db.get_options(q_id)   # (option_id, q_id, text, is_correct)
    option_vars  = []
    correct_var  = tk.IntVar(value=0)

    for i, opt in enumerate(options):
        if opt[3] == 1:   # opt[3] is is_correct
            correct_var.set(i)
        row = tk.Frame(form, bg=PANEL)
        row.pack(fill="x", pady=2)
        tk.Radiobutton(row, text=["A","B","C","D"][i], variable=correct_var, value=i,
                       bg=PANEL, fg=FG, selectcolor=ACCENT,
                       activebackground=PANEL).pack(side="left")
        var = tk.StringVar(value=opt[2])
        tk.Entry(row, textvariable=var, bg=ENTRY, fg=FG, relief="flat",
                 font=FONT_BODY, insertbackground=FG).pack(
                 side="left", fill="x", expand=True, ipady=5, padx=4)
        option_vars.append(var)

    error_var = tk.StringVar()
    tk.Label(form, textvariable=error_var, bg=PANEL, fg="#f87171",
             font=FONT_TINY).pack(pady=(4, 0))

    def save():
        new_text  = q_text_var.get().strip()
        new_marks = marks_var.get().strip()
        new_opts  = [v.get().strip() for v in option_vars]
        if not new_text:
            error_var.set("Question text is required."); return
        if not new_marks.isdigit() or int(new_marks) < 1:
            error_var.set("Enter valid marks."); return
        if any(o == "" for o in new_opts):
            error_var.set("All 4 options must be filled."); return

        db.update_question(q_id, new_text, int(new_marks))
        db.clear_correct_options(q_id)
        correct_index = correct_var.get()
        for i, opt in enumerate(options):
            is_correct = 1 if i == correct_index else 0
            db.update_option(opt[0], new_opts[i], is_correct)

        messagebox.showinfo("Saved", "Question updated!")
        on_done()
        win.destroy()

    make_btn(form, "Save Changes", save, color=GREEN).pack(pady=(12, 0), fill="x")


# ── All Attempts Window ────────────────────────────────────

def open_all_attempts_window(parent):
    """Table showing every student's every attempt."""
    win = tk.Toplevel(parent)
    win.title("All Attempts")
    win.configure(bg=BG)
    center_window(win, 640, 480)

    tk.Label(win, text="📊  All Student Attempts",
             font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(16, 4))

    # Column headers
    hdr = tk.Frame(win, bg="#3b3b55", pady=5)
    hdr.pack(fill="x", padx=20)
    for col_text, col_w in [("ID", 5), ("Student", 18), ("Quiz", 28), ("Score", 8), ("Status", 12)]:
        tk.Label(hdr, text=col_text, bg="#3b3b55", fg=ACCENT,
                 font=("Segoe UI", 10, "bold"), width=col_w,
                 anchor="w").pack(side="left", padx=4)

    # Attempt rows
    rows_frame = tk.Frame(win, bg=ENTRY)
    rows_frame.pack(fill="both", expand=True, padx=20, pady=4)

    scroll = tk.Scrollbar(rows_frame)
    scroll.pack(side="right", fill="y")

    attempt_box = tk.Listbox(rows_frame, bg=ENTRY, fg=FG, font=FONT_SMALL,
                             relief="flat", selectbackground=ACCENT,
                             activestyle="none", yscrollcommand=scroll.set)
    attempt_box.pack(fill="both", expand=True)
    scroll.config(command=attempt_box.yview)

    attempts = db.get_all_attempts()
    # Each row: (attempt_id, username, quiz_title, score, status)
    for a in attempts:
        attempt_box.insert(tk.END,
            f"  #{a[0]}   {a[1]:<18} {a[2]:<28} {a[3]:<8} {a[4]}")

    make_btn(win, "Close", win.destroy, color="#4b4b70").pack(pady=10)


# ── Student Progress Window ────────────────────────────────

def open_student_progress_window(parent):
    """
    Left panel: list of all students.
    Right panel: which quizzes they've attempted (and scores), which are pending.
    """
    win = tk.Toplevel(parent)
    win.title("Student Progress")
    win.configure(bg=BG)
    center_window(win, 800, 540)
    win.grab_set()

    tk.Label(win, text="👤  Student Progress",
             font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(16, 4))

    body = tk.Frame(win, bg=BG)
    body.pack(fill="both", expand=True, padx=16, pady=8)

    # ── Left: student list ────────────────────────────────
    left = tk.Frame(body, bg=BG, width=200)
    left.pack(side="left", fill="y", padx=(0, 12))
    left.pack_propagate(False)

    tk.Label(left, text="Students",
             font=("Segoe UI", 11, "bold"), bg=BG, fg=FG).pack(anchor="w")

    student_box = tk.Listbox(left, bg=ENTRY, fg=FG, font=FONT_BODY,
                             relief="flat", selectbackground=ACCENT,
                             activestyle="none")
    student_box.pack(fill="both", expand=True, pady=6)

    students = db.get_all_students()   # (user_id, username)
    for s in students:
        student_box.insert(tk.END, f"  {s[1]}")
    if not students:
        student_box.insert(tk.END, "  No students yet.")

    # ── Right: quiz progress table ────────────────────────
    right = tk.Frame(body, bg=BG)
    right.pack(side="left", fill="both", expand=True)

    header_label = tk.Label(right, text="Select a student to view progress",
                            font=("Segoe UI", 11, "bold"), bg=BG, fg=DIM)
    header_label.pack(anchor="w")

    # Stats bar
    stats_frame = tk.Frame(right, bg=PANEL, padx=12, pady=8)
    stats_frame.pack(fill="x", pady=(4, 6))

    attempted_var = tk.StringVar(value="Attempted:  —")
    pending_var   = tk.StringVar(value="Pending:    —")
    avg_var       = tk.StringVar(value="Avg Score:  —")

    for var in (attempted_var, pending_var, avg_var):
        tk.Label(stats_frame, textvariable=var, bg=PANEL, fg=FG,
                 font=FONT_SMALL).pack(side="left", padx=16)

    # Column headers for the quiz table
    col_hdr = tk.Frame(right, bg="#3b3b55", pady=5)
    col_hdr.pack(fill="x")
    for col_text, col_w in [("Quiz Title", 34), ("Time", 8),
                              ("Best Score", 12), ("Attempts", 10), ("Status", 14)]:
        tk.Label(col_hdr, text=col_text, bg="#3b3b55", fg=ACCENT,
                 font=("Segoe UI", 9, "bold"), width=col_w,
                 anchor="w").pack(side="left", padx=4)

    # Scrollable quiz rows
    table_frame = tk.Frame(right, bg=PANEL)
    table_frame.pack(fill="both", expand=True)

    table_canvas  = tk.Canvas(table_frame, bg=PANEL, highlightthickness=0)
    table_scroll  = tk.Scrollbar(table_frame, orient="vertical",
                                  command=table_canvas.yview)
    table_canvas.configure(yscrollcommand=table_scroll.set)
    table_scroll.pack(side="right", fill="y")
    table_canvas.pack(side="left", fill="both", expand=True)

    rows_inner = tk.Frame(table_canvas, bg=PANEL)
    canvas_win = table_canvas.create_window((0, 0), window=rows_inner, anchor="nw")

    def resize_canvas(e):
        table_canvas.itemconfig(canvas_win, width=e.width)
    table_canvas.bind("<Configure>", resize_canvas)

    def on_student_click(event):
        """When a student is selected, load and display their quiz progress."""
        sel = student_box.curselection()
        if not sel:
            return

        student_id, student_name = students[sel[0]]
        header_label.config(text=f"{student_name}  —  Quiz Progress", fg=FG)

        # progress rows: (quiz_id, title, time_limit, best_score, attempt_count)
        progress = db.get_student_progress(student_id)

        attempted = sum(1 for p in progress if p[4] > 0)
        pending   = sum(1 for p in progress if p[4] == 0)
        scores    = [p[3] for p in progress if p[3] is not None]
        avg       = round(sum(scores) / len(scores), 1) if scores else "N/A"

        attempted_var.set(f"Attempted: {attempted}")
        pending_var.set(  f"Pending:   {pending}")
        avg_var.set(      f"Avg Score: {avg}")

        # Clear old rows
        for widget in rows_inner.winfo_children():
            widget.destroy()

        for i, p in enumerate(progress):
            _, title, time_lim, best, count = p
            done     = (count > 0)
            row_bg   = ENTRY if i % 2 == 0 else PANEL
            status   = "Completed" if done else "Not Attempted"
            fg_score = GREEN if done else RED

            row = tk.Frame(rows_inner, bg=row_bg)
            row.pack(fill="x")

            tk.Label(row, text=f"  {title[:32]}", bg=row_bg, fg=FG,
                     font=FONT_SMALL, width=34, anchor="w").pack(side="left", padx=4, pady=3)
            tk.Label(row, text=f"{time_lim}m", bg=row_bg, fg=DIM,
                     font=FONT_SMALL, width=8, anchor="w").pack(side="left", padx=4)
            tk.Label(row, text=str(best) if done else "—", bg=row_bg, fg=fg_score,
                     font=FONT_SMALL, width=12, anchor="w").pack(side="left", padx=4)
            tk.Label(row, text=str(count), bg=row_bg, fg=DIM,
                     font=FONT_SMALL, width=10, anchor="w").pack(side="left", padx=4)
            tk.Label(row, text=status, bg=row_bg, fg=fg_score,
                     font=FONT_SMALL, width=14, anchor="w").pack(side="left", padx=4)

        rows_inner.update_idletasks()
        table_canvas.config(scrollregion=table_canvas.bbox("all"))

    student_box.bind("<<ListboxSelect>>", on_student_click)


# ── Import from Excel ──────────────────────────────────────

def open_excel_import_window(parent, on_done):
    """
    Lets the admin pick an Excel file and imports quizzes + questions in bulk.
    Expected columns: quiz_title, time_limit, question, marks,
                      option_a, option_b, option_c, option_d, correct (A/B/C/D)
    """
    if not EXCEL_OK:
        messagebox.showerror("Missing library",
                             "openpyxl is not installed.\n\nRun:\n  pip install openpyxl")
        return

    file_path = filedialog.askopenfilename(
        title="Choose Excel file",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
    )
    if not file_path:
        return

    # Read the workbook
    try:
        wb   = openpyxl.load_workbook(file_path)
        ws   = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except Exception as e:
        messagebox.showerror("Read Error", str(e))
        return

    # Process rows
    quiz_cache  = {}   # {quiz_title: quiz_id} — so we don't create duplicate quizzes
    added_count = 0
    errors      = []   # list of error strings for rows we couldn't import

    for row_num, row in enumerate(rows, start=2):
        if not any(row):         # skip completely empty rows
            continue
        if len(row) < 9:
            errors.append(f"Row {row_num}: Not enough columns (need 9).")
            continue

        quiz_title, time_limit, question, marks, op_a, op_b, op_c, op_d, correct = row[:9]

        # Validate
        if not all([quiz_title, time_limit, question, marks, op_a, op_b, op_c, op_d, correct]):
            errors.append(f"Row {row_num}: Missing data in one or more cells.")
            continue
        if str(correct).strip().upper() not in ("A", "B", "C", "D"):
            errors.append(f"Row {row_num}: 'correct' must be A, B, C, or D.")
            continue

        # Create or reuse the quiz
        quiz_title_str = str(quiz_title).strip()
        if quiz_title_str not in quiz_cache:
            quiz_id = db.add_quiz(quiz_title_str, int(time_limit))
            quiz_cache[quiz_title_str] = quiz_id
        else:
            quiz_id = quiz_cache[quiz_title_str]

        # Add the question
        q_id = db.add_question(quiz_id, str(question).strip(), int(marks))

        # Add options
        correct_letter = str(correct).strip().upper()
        for letter, opt_text in zip(["A","B","C","D"], [op_a, op_b, op_c, op_d]):
            is_correct = 1 if letter == correct_letter else 0
            db.add_option(q_id, str(opt_text).strip(), is_correct)

        added_count += 1

    on_done()   # refresh the quiz list in the dashboard

    # Show import report
    report_win = tk.Toplevel(parent)
    report_win.title("Import Report")
    report_win.configure(bg=BG)
    center_window(report_win, 460, 340)

    tk.Label(report_win, text="📥  Import Complete",
             font=FONT_TITLE, bg=BG, fg=FG).pack(pady=(20, 4))
    tk.Label(report_win, text=f"✅  {added_count} questions imported successfully.",
             bg=BG, fg=GREEN, font=FONT_BODY).pack(anchor="w", padx=20)

    if errors:
        tk.Label(report_win, text=f"⚠️  {len(errors)} rows skipped:",
                 bg=BG, fg=RED, font=FONT_BODY).pack(anchor="w", padx=20, pady=(8, 2))
        err_frame = tk.Frame(report_win, bg=PANEL)
        err_frame.pack(fill="both", expand=True, padx=20, pady=(4, 0))
        err_scroll = tk.Scrollbar(err_frame)
        err_scroll.pack(side="right", fill="y")
        err_box = tk.Listbox(err_frame, yscrollcommand=err_scroll.set,
                             bg=ENTRY, fg="#f87171", font=FONT_SMALL,
                             relief="flat", activestyle="none")
        err_box.pack(fill="both", expand=True)
        err_scroll.config(command=err_box.yview)
        for err in errors:
            err_box.insert(tk.END, f"  {err}")

    make_btn(report_win, "Close", report_win.destroy, color="#4b4b70").pack(pady=12)


# ── Export Gradebook to Excel ──────────────────────────────

def export_gradebook_to_excel(parent):
    """
    Saves all student attempt data to a styled .xlsx file.
    Creates one sheet per quiz + a combined 'All Attempts' summary sheet.
    """
    if not EXCEL_OK:
        messagebox.showerror("Missing library",
                             "openpyxl is not installed.\n\nRun:\n  pip install openpyxl")
        return

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import datetime

    save_path = filedialog.asksaveasfilename(
        title="Save Gradebook",
        defaultextension=".xlsx",
        initialfile=f"gradebook_{datetime.date.today()}.xlsx",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
    )
    if not save_path:
        return

    attempts = db.get_all_attempts()   # (attempt_id, username, quiz_title, score, status)
    if not attempts:
        messagebox.showinfo("No data", "No student attempts to export yet.")
        return

    wb = openpyxl.Workbook()

    # Styles
    header_fill = PatternFill("solid", fgColor="7c3aed")   # purple
    alt_fill    = PatternFill("solid", fgColor="2a2a3e")   # dark panel
    header_font = Font(bold=True, color="FFFFFF", size=11)
    bold_font   = Font(bold=True)
    center_align = Alignment(horizontal="center")

    def add_header_row(ws, columns):
        """Write a purple header row to the worksheet."""
        ws.append(columns)
        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center_align

    def auto_fit_columns(ws):
        """Set column widths based on content length."""
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # Sheet 1: all attempts summary
    ws_all = wb.active
    ws_all.title = "All Attempts"
    add_header_row(ws_all, ["#", "Student", "Quiz", "Score", "Status"])

    for i, a in enumerate(attempts):
        ws_all.append([a[0], a[1], a[2], a[3], a[4]])
        if i % 2 == 0:
            for cell in ws_all[i + 2]:
                cell.fill = alt_fill

    auto_fit_columns(ws_all)

    # One sheet per quiz
    quizzes = db.get_all_quizzes()
    for quiz in quizzes:
        quiz_id, title, _ = quiz
        # Sheet names: max 31 chars, no special characters
        sheet_name  = title[:28].replace("/", "-").replace("\\", "-")
        ws_q        = wb.create_sheet(title=sheet_name)
        add_header_row(ws_q, ["Student", "Score", "Status"])

        quiz_attempts = [a for a in attempts if a[2] == title]
        for i, a in enumerate(quiz_attempts):
            ws_q.append([a[1], a[3], a[4]])
            if i % 2 == 0:
                for cell in ws_q[i + 2]:
                    cell.fill = alt_fill

        # Stats rows at the bottom
        if quiz_attempts:
            scores = [a[3] for a in quiz_attempts]
            ws_q.append([])
            ws_q.append(["Total Attempts", len(quiz_attempts)])
            ws_q.append(["Average Score",  round(sum(scores) / len(scores), 1)])
            ws_q.append(["Highest Score",  max(scores)])
            ws_q.append(["Lowest Score",   min(scores)])
            for row in ws_q.iter_rows(min_row=ws_q.max_row - 3, max_row=ws_q.max_row):
                for cell in row:
                    cell.font = bold_font

        auto_fit_columns(ws_q)

    wb.save(save_path)
    messagebox.showinfo("Export complete", f"Gradebook saved!\n\n{save_path}")

    # Open the folder in Finder (macOS)
    import subprocess
    subprocess.Popen(["open", "-R", save_path])
